import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_run(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, color_rgb=(0, 0, 0), highlight=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    if highlight:
        hl = parse_xml(f'<w:highlight {nsdecls("w")} w:val="yellow"/>')
        run._r.get_or_add_rPr().append(hl)

def add_paragraph_formatted(doc, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, highlight=False):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    format_run(run, bold=bold, italic=italic, highlight=highlight)
    return p

def add_heading_formatted(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        format_run(r, bold=True)
    return h

def add_code_block(doc, title, code_lines):
    p = doc.add_paragraph()
    r = p.add_run(title)
    format_run(r, bold=True)
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    for idx, line in enumerate(code_lines):
        if idx == 0:
            cell_p = cell.paragraphs[0]
        else:
            cell_p = cell.add_paragraph()
        
        cell_p.paragraph_format.space_before = Pt(0)
        cell_p.paragraph_format.space_after = Pt(0)
        cell_p.paragraph_format.line_spacing = 1.15
        
        r = cell_p.add_run(line)
        format_run(r, font_name="Courier New", font_size=10)
    
    doc.add_paragraph()

def generate_exp3_docx():
    doc = Document()

    # Configure Normal Style default font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Title
    add_paragraph_formatted(doc, "EXPERIMENT 3", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_formatted(doc, "DEPLOYING AND INTERACTING WITH ERC-4626 VAULT CONTRACTS (GITBOUNTYS)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    # Aim
    add_heading_formatted(doc, "Aim", level=1)
    add_paragraph_formatted(
        doc,
        "To deploy and interact with an ERC-4626 tokenized vault using an ERC-20 token (MockUSDC) as the underlying asset, "
        "and to demonstrate the complete asset-to-vault-share and vault-share-to-asset lifecycle within the GitBountys developer platform."
    )
    doc.add_paragraph()

    # Scope
    add_heading_formatted(doc, "Scope of the Experiment", level=1)
    add_paragraph_formatted(
        doc,
        "This experiment is restricted to ERC-20 and ERC-4626 vault concepts. Do not implement liquidity pools, "
        "AMM, yield farming, staking, or a DEX in this experiment; those mechanisms belong to subsequent experiments.\n"
        "• Deploy a mock ERC-20 token to act as the underlying asset.\n"
        "• Deploy an ERC-4626 vault (BountyVault) that accepts the mock ERC-20 asset.\n"
        "• Interact with the ERC-20 standard functions required for the vault workflow.\n"
        "• Interact with the ERC-4626 standardized vault functions.\n"
        "• Add one simple application-specific vault feature (harvestYield sponsor bonus deposit) without breaking standard vault behaviour.\n"
        "• Demonstrate approval, deposit, receipt of vault shares, conversion, withdrawal/redeem, and returned underlying assets."
    )
    doc.add_paragraph()

    # Standards Table
    add_heading_formatted(doc, "Standards to be Demonstrated", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Layer", "Functions / Behaviour", "Purpose"]
    
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "E6E6E6")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        for r in p.runs:
            format_run(r, bold=True)

    rows_data = [
        ("ERC-20 underlying asset", "totalSupply(), balanceOf(), transfer(), approve(), allowance(), transferFrom(), Transfer/Approval events", "Provide and authorize the bounty asset used by the vault."),
        ("ERC-4626 vault", "asset(), totalAssets(), deposit(), mint(), withdraw(), redeem(), convertToShares(), convertToAssets(), preview...(), max...(), Deposit/Withdraw events", "Standardize asset/share accounting and vault entry/exit for developers."),
        ("ERC-20-compatible vault shares", "balanceOf(), totalSupply(), transfer(), approve(), allowance(), transferFrom()", "Represent developer ownership/claim on the GitBountys vault."),
        ("Application-specific logic", "harvestYield(uint256 yieldAmount)", "Demonstrate GitBountys sponsor yield/bonus injection without breaking standard vault behaviour.")
    ]

    for layer, funcs, purpose in rows_data:
        row_cells = table.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = funcs
        row_cells[2].text = purpose
        for cell in row_cells:
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            for r in p.runs:
                format_run(r)

    doc.add_paragraph()

    # Smart Contracts Source Code
    add_heading_formatted(doc, "Smart Contracts Source Code", level=1)

    vault_code = [
        "// SPDX-License-Identifier: MIT",
        "pragma solidity ^0.8.24;",
        "",
        "import \"@openzeppelin/contracts/token/ERC20/extensions/ERC4626.sol\";",
        "import \"@openzeppelin/contracts/token/ERC20/ERC20.sol\";",
        "",
        "contract BountyVault is ERC4626 {",
        "    uint256 public totalYieldHarvested;",
        "",
        "    event YieldHarvested(address indexed harvester, uint256 yieldAmount, uint256 newTotalAssets);",
        "",
        "    constructor(IERC20 asset_)",
        "        ERC20(\"GitBountys Vault Share\", \"vUSDC\")",
        "        ERC4626(asset_)",
        "    {}",
        "",
        "    function harvestYield(uint256 yieldAmount) external {",
        "        require(yieldAmount > 0, \"Yield amount must be greater than 0\");",
        "        IERC20(asset()).transferFrom(msg.sender, address(this), yieldAmount);",
        "        totalYieldHarvested += yieldAmount;",
        "        emit YieldHarvested(msg.sender, yieldAmount, totalAssets());",
        "    }",
        "}"
    ]
    add_code_block(doc, "BountyVault.sol (GitBountys ERC-4626 Vault with Sponsor Yield Feature):", vault_code)

    mock_usdc_code = [
        "// SPDX-License-Identifier: MIT",
        "pragma solidity ^0.8.24;",
        "",
        "import \"@openzeppelin/contracts/token/ERC20/ERC20.sol\";",
        "",
        "contract MockERC20 is ERC20 {",
        "    uint8 private immutable _customDecimals;",
        "",
        "    constructor(string memory name_, string memory symbol_, uint256 initialSupply, uint8 decimals_)",
        "        ERC20(name_, symbol_)",
        "    {",
        "        _customDecimals = decimals_;",
        "        _mint(msg.sender, initialSupply);",
        "    }",
        "",
        "    function decimals() public view virtual override returns (uint8) {",
        "        return _customDecimals;",
        "    }",
        "",
        "    function mint(address to, uint256 amount) external {",
        "        _mint(to, amount);",
        "    }",
        "}"
    ]
    add_code_block(doc, "MockERC20.sol (Underlying ERC-20 Bounty Asset):", mock_usdc_code)

    # Interaction Flow Steps
    add_heading_formatted(doc, "Expected Interaction Flow & Verification", level=1)
    flow_steps = [
        "1. Deploy MockUSDC underlying ERC-20 contract.",
        "2. Mint test MockUSDC tokens to developer wallet (e.g. 1000 USDC).",
        "3. Deploy BountyVault passing MockUSDC address to constructor.",
        "4. Verify asset() returns MockUSDC contract address.",
        "5. Approve vault: approve(BountyVault, 500 USDC) on MockUSDC.",
        "6. Verify allowance(wallet, BountyVault) == 500 USDC.",
        "7. Call previewDeposit(100 USDC) and convertToShares(100 USDC).",
        "8. Execute deposit(100 USDC, receiver) -> receive vUSDC vault shares.",
        "9. Verify post-deposit state: MockUSDC balance decreased, vUSDC balance increased, totalAssets() increased.",
        "10. Execute convertToShares() and convertToAssets() conversion tests.",
        "11. Execute mint(50 vUSDC, receiver) -> Observe required underlying asset amount transferred.",
        "12. Execute withdraw(30 USDC, receiver, owner) -> Underlying asset returned.",
        "13. Execute redeem(40 vUSDC, receiver, owner) -> Shares burned, underlying asset returned.",
        "14. Demonstrate GitBountys custom feature: harvestYield(200 USDC) -> Inject sponsor yield into vault without issuing shares -> Observe totalAssets() increase and share exchange rate appreciation for developers.",
        "15. Audit on-chain events: Transfer, Approval, Deposit, Withdraw, YieldHarvested."
    ]

    for step in flow_steps:
        add_paragraph_formatted(doc, step)

    doc.add_paragraph()

    # Viva Questions & Answers
    add_heading_formatted(doc, "Viva Questions & Technical Answers", level=1)
    viva_qna = [
        ("1. What is the underlying asset in your vault?", "The underlying asset is MockUSDC (an ERC-20 token contract representing bounty rewards)."),
        ("2. What does an ERC-4626 vault share represent?", "An ERC-4626 vault share (vUSDC) represents a pro-rata fractional ownership claim on the total underlying assets managed by the GitBountys vault."),
        ("3. Why is an ERC-20 approval required before deposit?", "ERC-4626 deposit() calls transferFrom() to move underlying tokens from the user's wallet to the vault. The user must first approve the vault contract allowance."),
        ("4. What is the difference between deposit() and mint()?", "deposit() accepts a specified amount of underlying assets and calculates shares to issue. mint() accepts a specified amount of vault shares to issue and calculates required underlying assets to deposit."),
        ("5. What is the difference between withdraw() and redeem()?", "withdraw() takes a target underlying asset amount to return and calculates required shares to burn. redeem() takes a target vault share amount to burn and calculates underlying assets to return."),
        ("6. What does totalAssets() represent?", "totalAssets() returns the total quantity of underlying ERC-20 tokens currently held and managed by the vault contract."),
        ("7. What is the purpose of convertToShares() and convertToAssets()?", "They perform on-chain rate conversions between underlying assets and vault shares based on current totalAssets() and totalSupply()."),
        ("8. Why are ERC-4626 vault shares ERC-20 compatible?", "ERC-4626 inherits ERC-20. Vault shares possess standard ERC-20 functions (transfer, balanceOf, approve) allowing developer share tokens to be traded or integrated."),
        ("9. Which functions did you inherit and which function did you add?", "Inherited standard OpenZeppelin ERC4626 & ERC20 functions. Added custom application feature harvestYield(uint256 yieldAmount)."),
        ("10. What happens to the user's underlying ERC-20 balance after deposit and after redeem?", "After deposit, underlying ERC-20 balance decreases while vault share balance increases. After redeem, vault shares are burned and underlying ERC-20 balance increases.")
    ]

    for q, a in viva_qna:
        add_paragraph_formatted(doc, q, bold=True)
        add_paragraph_formatted(doc, a, italic=True)
        doc.add_paragraph()

    out_file = os.path.join(os.path.dirname(__file__), 'Defi Exp 3 Vaults.docx')
    doc.save(out_file)
    print(f"Successfully generated Experiment 3 report at: {out_file}")

if __name__ == '__main__':
    generate_exp3_docx()
